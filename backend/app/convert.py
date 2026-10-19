"""Features that produce new files: export, extract, split, combine, create, compare, check."""

from __future__ import annotations

import csv
import difflib
import io
import json
import re
import tempfile
import zipfile
from pathlib import Path

import pymupdf as fitz

from .ops.pages import open_any_as_pdf
from .store import DocEntry, save_bytes

# ---------------------------------------------------------------- export --


def export_docx(entry: DocEntry) -> bytes:
    from pdf2docx import Converter

    with tempfile.TemporaryDirectory() as tmp:
        src, dst = Path(tmp, "in.pdf"), Path(tmp, "out.docx")
        src.write_bytes(entry.plain_bytes())
        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        finally:
            cv.close()
        return dst.read_bytes()


def export_xlsx(entry: DocEntry) -> bytes:
    """One worksheet per detected table; a text sheet per page when no table is found."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    doc = entry.open()
    try:
        for page in doc:
            tables = page.find_tables().tables
            for t_index, table in enumerate(tables, 1):
                ws = wb.create_sheet(f"P{page.number + 1} T{t_index}"[:31])
                for row in table.extract():
                    ws.append([c if c is not None else "" for c in row])
            if not tables:
                ws = wb.create_sheet(f"P{page.number + 1} text"[:31])
                for line in page.get_text("text").splitlines():
                    ws.append([line])
    finally:
        doc.close()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pptx(entry: DocEntry, dpi: int = 150) -> bytes:
    """One slide per page, each page as a full-slide image."""
    from pptx import Presentation
    from pptx.util import Emu

    doc = entry.open()
    try:
        first = doc[0].rect
        prs = Presentation()
        prs.slide_width = Emu(int(first.width * 12700))
        prs.slide_height = Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        for page in doc:
            slide = prs.slides.add_slide(blank)
            png = page.get_pixmap(dpi=dpi).tobytes("png")
            slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
            notes = page.get_text("text").strip()
            if notes:
                slide.notes_slide.notes_text_frame.text = notes[:5000]
    finally:
        doc.close()
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export_images(entry: DocEntry, fmt: str = "png", dpi: int = 150,
                  pages: list[int] | None = None) -> bytes:
    if fmt not in ("png", "jpg"):
        raise ValueError("Format must be png or jpg")
    doc = entry.open()
    buf = io.BytesIO()
    try:
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for i in pages if pages is not None else range(doc.page_count):
                pix = doc[i].get_pixmap(dpi=dpi)
                data = pix.tobytes("png") if fmt == "png" else pix.tobytes("jpg", jpg_quality=90)
                zf.writestr(f"page-{i + 1:04d}.{fmt}", data)
    finally:
        doc.close()
    return buf.getvalue()


def export_text(entry: DocEntry, fmt: str = "txt") -> bytes:
    doc = entry.open()
    try:
        if fmt == "txt":
            return "\n\f".join(p.get_text("text") for p in doc).encode()
        if fmt == "html":
            body = "\n".join(p.get_text("html") for p in doc)
            return f"<!doctype html><meta charset='utf-8'><title>{entry.name}</title>{body}".encode()
        if fmt == "md":
            return _markdown(doc).encode()
        raise ValueError("Format must be txt, html or md")
    finally:
        doc.close()


def _markdown(doc: fitz.Document) -> str:
    sizes: dict[float, int] = {}
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                for s in ln["spans"]:
                    sizes[round(s["size"])] = sizes.get(round(s["size"]), 0) + len(s["text"])
    body = max(sizes, key=sizes.get) if sizes else 11
    out = []
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            lines = b.get("lines", [])
            text = " ".join("".join(s["text"] for s in ln["spans"]) for ln in lines).strip()
            if not text:
                continue
            size = max((s["size"] for ln in lines for s in ln["spans"]), default=body)
            if size >= body * 1.6:
                out.append(f"# {text}")
            elif size >= body * 1.3:
                out.append(f"## {text}")
            elif size >= body * 1.15:
                out.append(f"### {text}")
            else:
                out.append(text)
        out.append("\n---\n")
    return "\n\n".join(out)


def export_form_data(entry: DocEntry, fmt: str = "json") -> bytes:
    doc = entry.open()
    try:
        values = {w.field_name: w.field_value for p in doc for w in p.widgets()}
    finally:
        doc.close()
    if fmt == "json":
        return json.dumps(values, indent=2, default=str).encode()
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(list(values))
        writer.writerow([values[k] for k in values])
        return buf.getvalue().encode()
    raise ValueError("Format must be json or csv")


def export_comments(entry: DocEntry, fmt: str = "csv") -> bytes:
    doc = entry.open()
    try:
        rows = []
        for page in doc:
            for a in page.annots():
                if a.type[1] in ("Popup", "Widget", "Link"):
                    continue
                rows.append({"page": page.number + 1, "type": a.type[1],
                             "author": a.info.get("title", ""), "contents": a.info.get("content", ""),
                             "modified": a.info.get("modDate", ""),
                             "marked_text": page.get_textbox(a.rect).strip()
                             if a.type[1] in ("Highlight", "Underline", "StrikeOut", "Squiggly") else ""})
    finally:
        doc.close()
    if fmt == "json":
        return json.dumps(rows, indent=2).encode()
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=["page", "type", "author", "contents", "modified", "marked_text"])
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode()


# ------------------------------------------------------ extract / split --

def parse_ranges(spec: str, page_count: int) -> list[list[int]]:
    """'1-3, 5, 8-' -> [[0,1,2],[4],[7..last]] (1-based input)."""
    groups = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d*)\s*-\s*(\d*)|(\d+)", part)
        if not m:
            raise ValueError(f"Bad page range {part!r}")
        if m.group(3):
            a = b = int(m.group(3))
        else:
            a = int(m.group(1) or 1)
            b = int(m.group(2) or page_count)
        if not 1 <= a <= b <= page_count:
            raise ValueError(f"Page range {part!r} is outside 1-{page_count}")
        groups.append(list(range(a - 1, b)))
    if not groups:
        raise ValueError("No pages selected")
    return groups


def extract_pages(entry: DocEntry, pages: list[int]) -> bytes:
    doc = entry.open()
    try:
        bad = [p for p in pages if not 0 <= p < doc.page_count]
        if bad or not pages:
            raise ValueError("Invalid page selection")
        out = fitz.open()
        for p in pages:
            out.insert_pdf(doc, from_page=p, to_page=p)
        return save_bytes(out)
    finally:
        doc.close()


def split(entry: DocEntry, mode: str, every: int = 1, ranges: str = "",
          max_mb: float = 0) -> bytes:
    doc = entry.open()
    try:
        n = doc.page_count
        if mode == "every":
            if every < 1:
                raise ValueError("every must be at least 1")
            groups = [list(range(i, min(i + every, n))) for i in range(0, n, every)]
        elif mode == "ranges":
            groups = parse_ranges(ranges, n)
        elif mode == "bookmarks":
            starts = sorted({p - 1 for lvl, _t, p in doc.get_toc() if lvl == 1 and p >= 1})
            if not starts:
                raise ValueError("The document has no top-level bookmarks")
            if starts[0] != 0:
                starts.insert(0, 0)
            groups = [list(range(s, e)) for s, e in zip(starts, starts[1:] + [n])]
        elif mode == "size":
            if max_mb <= 0:
                raise ValueError("max_mb must be positive")
            groups, current = [], []
            for i in range(n):
                trial = current + [i]
                probe = fitz.open()
                for p in trial:
                    probe.insert_pdf(doc, from_page=p, to_page=p)
                if current and len(save_bytes(probe)) > max_mb * 1024 * 1024:
                    groups.append(current)
                    current = [i]
                else:
                    current = trial
            if current:
                groups.append(current)
        else:
            raise ValueError("mode must be every, ranges, bookmarks or size")
        stem = entry.name.rsplit(".", 1)[0]
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for k, group in enumerate(groups, 1):
                part = fitz.open()
                for p in group:
                    part.insert_pdf(doc, from_page=p, to_page=p)
                zf.writestr(f"{stem}-part{k:03d}.pdf", save_bytes(part))
        return buf.getvalue()
    finally:
        doc.close()


# ------------------------------------------------------ combine / create --

def combine(parts: list[tuple[str, bytes]], bookmarks: bool = True) -> bytes:
    if len(parts) < 1:
        raise ValueError("Nothing to combine")
    out = fitz.open()
    toc = []
    for name, data in parts:
        src = open_any_as_pdf(name, data)
        start = out.page_count
        out.insert_pdf(src)
        if bookmarks:
            toc.append([1, name.rsplit(".", 1)[0], start + 1])
            toc.extend([lvl + 1, t, p + start] for lvl, t, p in src.get_toc())
    if bookmarks:
        out.set_toc(toc)
    return save_bytes(out)


def create_from(name: str, data: bytes) -> bytes:
    """Create a PDF from an image, text, Markdown, HTML or any MuPDF-readable file."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in ("txt", "md", "markdown", "html", "htm"):
        text = data.decode("utf-8", errors="replace")
        html = text if ext in ("html", "htm") else _text_to_html(text, markdown=ext != "txt")
        return html_to_pdf(html)
    return save_bytes(open_any_as_pdf(name, data))


def _text_to_html(text: str, markdown: bool) -> str:
    import html as h

    if not markdown:
        return f"<pre style='font-family: sans-serif; white-space: pre-wrap'>{h.escape(text)}</pre>"
    out = []
    for block in re.split(r"\n\s*\n", text):
        b = block.strip()
        m = re.match(r"(#{1,6})\s+(.*)", b)
        if m:
            lvl = len(m.group(1))
            out.append(f"<h{lvl}>{h.escape(m.group(2))}</h{lvl}>")
        elif all(re.match(r"[-*]\s+", ln) for ln in b.splitlines()):
            items = "".join(f"<li>{h.escape(ln[2:].strip())}</li>" for ln in b.splitlines())
            out.append(f"<ul>{items}</ul>")
        else:
            inner = h.escape(b).replace("\n", "<br>")
            inner = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", inner)
            inner = re.sub(r"\*(.+?)\*", r"<i>\1</i>", inner)
            out.append(f"<p>{inner}</p>")
    return "\n".join(out)


def html_to_pdf(html: str) -> bytes:
    story = fitz.Story(html=html, user_css="body {font-family: sans-serif; font-size: 11pt;}")
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    mediabox = fitz.paper_rect("letter")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buf.getvalue()


def blank_pdf(pages: int = 1, size: str = "letter") -> bytes:
    out = fitz.open()
    rect = fitz.paper_rect(size)
    for _ in range(max(1, pages)):
        out.new_page(width=rect.width, height=rect.height)
    return save_bytes(out)


# ---------------------------------------------------------------- compare --

def _words(doc: fitz.Document) -> list[tuple[int, fitz.Rect, str]]:
    out = []
    for page in doc:
        for w in page.get_text("words", sort=True):
            out.append((page.number, fitz.Rect(w[:4]), w[4]))
    return out


def compare(old: DocEntry, new: DocEntry) -> tuple[dict, bytes]:
    """Word-level diff of two documents. Returns a report and an annotated copy of `new`."""
    a_doc, b_doc = old.open(), new.open()
    try:
        a, b = _words(a_doc), _words(b_doc)
        sm = difflib.SequenceMatcher(a=[w[2] for w in a], b=[w[2] for w in b], autojunk=False)
        out = fitz.open(stream=new.plain_bytes(), filetype="pdf")
        changes = []
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "equal":
                continue
            old_text = " ".join(w[2] for w in a[i1:i2])
            new_text = " ".join(w[2] for w in b[j1:j2])
            if j2 > j1:
                by_page: dict[int, list[fitz.Rect]] = {}
                for p, r, _t in b[j1:j2]:
                    by_page.setdefault(p, []).append(r)
                for p, rects in by_page.items():
                    target = out[p]  # keep the page alive while its annotation is edited
                    annot = target.add_highlight_annot([r.quad for r in rects])
                    color = (0.55, 0.85, 0.55) if tag == "insert" else (1, 0.8, 0.3)
                    annot.set_colors(stroke=color)
                    label = "Inserted" if tag == "insert" else f"Replaced: {old_text[:300]}"
                    annot.set_info(title="Compare", content=label)
                    annot.update()
                page = b[j1][0]
            else:
                anchor = b[j1] if j1 < len(b) else (b[-1] if b else None)
                page = anchor[0] if anchor else 0
                if anchor:
                    r = anchor[1]
                    target = out[page]
                    note = target.add_text_annot(fitz.Point(max(r.x0 - 14, 0), r.y0),
                                                 f"Deleted: {old_text[:500]}", icon="Note")
                    note.set_colors(stroke=(0.9, 0.2, 0.2))
                    note.set_info(title="Compare", content=f"Deleted: {old_text[:500]}")
                    note.update()
            changes.append({"type": {"insert": "inserted", "delete": "deleted",
                                     "replace": "replaced"}[tag],
                            "page": page, "old": old_text[:500], "new": new_text[:500]})
        report = {
            "old": old.name, "new": new.name,
            "old_pages": a_doc.page_count, "new_pages": b_doc.page_count,
            "similarity": round(sm.ratio() * 100, 1),
            "counts": {k: sum(1 for c in changes if c["type"] == k)
                       for k in ("inserted", "deleted", "replaced")},
            "changes": changes,
        }
        return report, save_bytes(out)
    finally:
        a_doc.close()
        b_doc.close()


# ---------------------------------------------------------- accessibility --

def accessibility_check(entry: DocEntry) -> dict:
    doc = entry.open()
    try:
        cat = doc.pdf_catalog()
        checks = []

        def add(rule: str, passed: bool | None, detail: str) -> None:
            checks.append({"rule": rule, "status": "passed" if passed else (
                "manual" if passed is None else "failed"), "detail": detail})

        marked = doc.xref_get_key(cat, "MarkInfo/Marked")[1] == "true"
        struct = doc.xref_get_key(cat, "StructTreeRoot")[0] != "null"
        add("Tagged PDF", marked and struct, "Document structure tags present" if marked and struct
            else "No structure tags; screen readers cannot follow the reading order")
        lang = doc.xref_get_key(cat, "Lang")
        add("Primary language", lang[0] != "null", f"Lang = {lang[1]}" if lang[0] != "null"
            else "Document language is not set")
        title = (doc.metadata or {}).get("title", "")
        add("Title", bool(title), f"Title = {title!r}" if title else "Document title is empty")
        display = doc.xref_get_key(cat, "ViewerPreferences/DisplayDocTitle")[1] == "true"
        add("Display title in window", display, "DisplayDocTitle set" if display
            else "Viewer shows file name instead of title")
        image_only = [p.number + 1 for p in doc if not p.get_text("text").strip() and p.get_images()]
        add("Text is real text (not scanned images)", not image_only,
            "All pages have text" if not image_only else f"Image-only pages: {image_only[:20]} (run OCR)")
        unembedded = sorted({f[3] for p in doc for f in p.get_fonts() if f[1] in ("n/a", "")
                             and not f[3].startswith(("Helvetica", "Times", "Courier", "Symbol", "Zapf"))})
        add("Fonts embedded", not unembedded, "All fonts embedded" if not unembedded
            else f"Not embedded: {', '.join(unembedded[:10])}")
        add("Bookmarks for long documents", doc.page_count <= 20 or bool(doc.get_toc()),
            f"{len(doc.get_toc())} bookmarks for {doc.page_count} pages")
        no_tip = [w.field_name for p in doc for w in p.widgets() if not w.field_label]
        add("Form fields have descriptions", not no_tip, "All fields have tooltips" if not no_tip
            else f"Fields without tooltip: {', '.join(no_tip[:10])}")
        perms_ok = bool(doc.permissions & fitz.PDF_PERM_ACCESSIBILITY) or not doc.is_encrypted
        add("Security allows assistive technology", perms_ok,
            "Allowed" if perms_ok else "Encryption blocks screen readers")
        add("Alternate text for images", None if struct else False,
            "Verify each figure has alt text" if struct else "Untagged: images cannot have alt text")
        add("Color contrast", None, "Check manually")
        failed = sum(c["status"] == "failed" for c in checks)
        return {"passed": sum(c["status"] == "passed" for c in checks), "failed": failed,
                "manual": sum(c["status"] == "manual" for c in checks), "checks": checks}
    finally:
        doc.close()
